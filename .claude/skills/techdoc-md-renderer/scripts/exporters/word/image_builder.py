"""Word image builder."""

from __future__ import annotations

import base64
import os
import tempfile
from pathlib import Path

from bs4 import Tag
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

try:
    import requests
    HAS_REQUESTS = True
except ImportError:  # pragma: no cover
    requests = None  # type: ignore
    HAS_REQUESTS = False


class ImageBuilder:
    """Render HTML figure/image nodes into Word."""

    MAX_IMAGE_W_EMU = int(12 * 360000)
    IMG_DPI_FALLBACK = 96

    def add_image(self, tag: Tag, doc):
        img_tag = tag.find('img')
        if not img_tag:
            return
        src = img_tag.get('src', '')
        alt = img_tag.get('alt', '')
        caption = tag.find('figcaption')
        caption_text = caption.get_text() if caption else alt

        if not src:
            self._add_placeholder(doc, caption_text)
            return

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        temp_path = None

        try:
            img_w, img_h = (None, None)
            image_arg = src

            if src.startswith('data:'):
                header, encoded = src.split(',', 1)
                ext = header.split(';')[0].split('/')[-1] if 'image/' in header else 'png'
                data = base64.b64decode(encoded)
                img_w, img_h = self._get_image_physical_size(data)
                with tempfile.NamedTemporaryFile(suffix=f'.{ext}', delete=False) as f:
                    f.write(data)
                    image_arg = f.name
                    temp_path = f.name
            elif src.startswith(('http://', 'https://')):
                if HAS_REQUESTS:
                    try:
                        resp = requests.get(src, timeout=10)
                        resp.raise_for_status()
                        img_w, img_h = self._get_image_physical_size(resp.content)
                    except Exception:
                        pass
                image_arg = src
            elif Path(src).exists():
                img_w, img_h = self._get_image_physical_size_file(src)
                image_arg = src
            else:
                self._add_placeholder_run(para, caption_text or src)
                return

            if img_w and img_h:
                if img_w > self.MAX_IMAGE_W_EMU:
                    ratio = self.MAX_IMAGE_W_EMU / img_w
                    img_w = self.MAX_IMAGE_W_EMU
                    img_h = int(img_h * ratio)
                para.add_run().add_picture(image_arg, width=img_w, height=img_h)
            else:
                para.add_run().add_picture(image_arg, width=self.MAX_IMAGE_W_EMU)
        except Exception:
            self._add_placeholder_run(para, caption_text or src)
        finally:
            if temp_path:
                try:
                    os.unlink(temp_path)
                except OSError:
                    pass

        if caption_text:
            cap_para = doc.add_paragraph()
            cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap_para.add_run(f'图 {caption_text}')
            cap_run.font.size = Pt(9)
            cap_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    def add_flowchart_element(self, tag: Tag, doc):
        if tag.find('img'):
            self.add_image(tag, doc)

    @staticmethod
    def _add_placeholder(doc, text: str):
        if text:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ImageBuilder._add_placeholder_run(para, text)

    @staticmethod
    def _add_placeholder_run(para, text: str):
        run = para.add_run(f'[图片: {text}]')
        run.font.italic = True

    @classmethod
    def _get_image_physical_size(cls, data: bytes):
        from io import BytesIO
        from PIL import Image
        img = Image.open(BytesIO(data))
        dpi = img.info.get('dpi', (cls.IMG_DPI_FALLBACK, cls.IMG_DPI_FALLBACK))
        dpi_x = dpi[0] if dpi[0] and dpi[0] > 0 else cls.IMG_DPI_FALLBACK
        dpi_y = dpi[1] if dpi[1] and dpi[1] > 0 else cls.IMG_DPI_FALLBACK
        return int(img.width / dpi_x * 914400), int(img.height / dpi_y * 914400)

    @classmethod
    def _get_image_physical_size_file(cls, path: str):
        with open(path, 'rb') as f:
            return cls._get_image_physical_size(f.read())

