"""HTML 表格 → Word 表格构建器

将 BeautifulSoup4 解析的 HTML <table> 转换为 python-docx 表格对象。
处理表头灰底、边框、列对齐、单元格边距等样式。
字体从文档 Normal 样式继承，确保与正文一致。
"""

from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup, NavigableString
from pathlib import Path
import base64
import hashlib
import os
import tempfile

from analyzers.table_classifier import TableClassifier
from .style_mapper import StyleMapper

try:
    import requests
except ImportError:
    requests = None

# 表格样式常量（字体从 Normal 继承，此处仅定义非字体属性）
HEADER_BG = 'D9D9D9'
BORDER_COLOR = '808080'
BORDER_SIZE = 4
CELL_MARGIN = 40


def _get_normal_font(doc) -> dict:
    """从文档 Normal 样式提取字体配置，用于表格/代码块等"""
    try:
        ns = doc.styles['Normal']
    except KeyError:
        return {'west': 'Consolas', 'east': 'Microsoft YaHei', 'size': Pt(10)}
    rPr = ns.element.rPr if ns.element.rPr is not None else None
    rFonts = rPr.find(qn('w:rFonts')) if rPr is not None else None
    west = rFonts.get(qn('w:ascii')) if rFonts is not None else None
    east = rFonts.get(qn('w:eastAsia')) if rFonts is not None else None
    size = ns.font.size
    return {
        'west': west or ns.font.name or 'Consolas',
        'east': east or west or ns.font.name or 'Microsoft YaHei',
        'size': size or Pt(10),
    }


class TableBuilder:
    """将 HTML <table> 标签转换为 python-docx 表格"""

    def __init__(self, exporter):
        """
        Args:
            exporter: WordExporter 实例，提供 _has_style / _ensure_element 工具方法
        """
        self._exp = exporter
        self._doc = None

    def build(self, tag, doc):
        """从 HTML table 标签构建 Word 表格"""
        self._doc = doc
        # 收集所有行
        all_rows = self._collect_rows(tag)
        if not all_rows:
            return

        col_count = max(len(r) for r in all_rows)
        has_thead = tag.find('thead') is not None
        explicit_kind = tag.get('data-table-kind') or ''
        analysis = TableClassifier.classify([
            [cell.get('text', '') for cell in row]
            for row in all_rows
        ], explicit_kind=explicit_kind)

        if analysis.layout.landscape:
            self._begin_landscape_section(doc)

        table = doc.add_table(rows=len(all_rows), cols=col_count)
        table.autofit = False
        table_style_name = StyleMapper.get_table_style(tag)
        if self._exp._has_style(doc, table_style_name):
            table.style = doc.styles[table_style_name]

        self._set_borders(table, analysis.layout.widths)

        # 填充数据
        header_count = len(self._collect_header_rows(tag))
        for row_idx, row_data in enumerate(all_rows):
            row = table.rows[row_idx]
            is_header_row = row_idx < header_count
            if not has_thead and row_idx == 0:
                first_tag = tag.find_all('tr')
                if row_idx < len(first_tag):
                    is_header_row = bool(first_tag[row_idx].find('th'))

            for col_idx, cell_data in enumerate(row_data):
                if col_idx >= col_count:
                    break
                cell = row.cells[col_idx]
                self._format_cell(
                    cell,
                    cell_data,
                    is_header_row,
                    font_size_pt=analysis.layout.font_size_pt,
                    force_center=col_idx in analysis.layout.center_columns,
                    code_font=col_idx in analysis.layout.code_columns,
                )

        self._apply_layout(table, analysis.layout.widths, header_count if has_thead else 0)

        if analysis.layout.landscape:
            self._end_portrait_section(doc)

    @staticmethod
    def _collect_header_rows(tag):
        """收集 thead 中的行"""
        rows = []
        thead = tag.find('thead')
        if thead:
            for tr in thead.find_all('tr', recursive=False):
                cells = list(StyleMapper.cell_collector(tr, True))
                if cells:
                    rows.append(cells)
        return rows

    @staticmethod
    def _collect_rows(tag):
        """收集所有行（thead + tbody）"""
        header_rows = TableBuilder._collect_header_rows(tag)
        has_thead = tag.find('thead') is not None

        tbody = tag.find('tbody') or tag
        body_rows = []
        for tr in tbody.find_all('tr', recursive=False):
            if has_thead and tr.parent and tr.parent.name == 'thead':
                continue
            is_header = (not has_thead and len(body_rows) == 0 and
                         (tr.find('th') or StyleMapper.is_thead_row(tr)))
            cells = list(StyleMapper.cell_collector(tr, is_header if not has_thead else False))
            if cells:
                body_rows.append(cells)

        return header_rows + body_rows

    def _format_cell(self, cell, cell_data: dict, is_header: bool,
                     font_size_pt=None, force_center: bool = False,
                     code_font: bool = False):
        """格式化单个表格单元格，字体从文档 Normal 样式继承"""
        text = cell_data.get('text', '')
        element = cell_data.get('element')
        align_cls = cell_data.get('align')

        tc = cell._tc
        tcPr = self._exp._ensure_element(tc, 'w:tcPr', first=True)

        # 边距
        old_mar = tcPr.find(qn('w:tcMar'))
        if old_mar is not None:
            tcPr.remove(old_mar)
        tcMar = OxmlElement('w:tcMar')
        for side, val in [('top', CELL_MARGIN), ('left', CELL_MARGIN + 20),
                          ('bottom', CELL_MARGIN), ('right', CELL_MARGIN + 20)]:
            m = OxmlElement(f'w:{side}')
            m.set(qn('w:w'), str(val))
            m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)

        # 垂直居中
        vAlign = self._exp._ensure_element(tcPr, 'w:vAlign')
        vAlign.set(qn('w:val'), 'center')

        # 表头灰底
        if is_header:
            shd = self._exp._ensure_element(tcPr, 'w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), HEADER_BG)

        # 获取文档 Normal 样式的字体配置
        normal_font = _get_normal_font(self._doc)

        # 写入文本
        para = cell.paragraphs[0]
        para.clear()
        if code_font and len(text.strip()) > 120:
            self._write_multiline_text(cell, self._wrap_long_code_cell(text))
        elif element is not None:
            self._render_cell_element(cell, element, text)
        else:
            self._write_multiline_text(cell, text)

        # 对齐
        if is_header or force_center or align_cls == 'align-center':
            align = WD_ALIGN_PARAGRAPH.CENTER
        elif align_cls == 'align-right':
            align = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            align = WD_ALIGN_PARAGRAPH.LEFT

        for p in cell.paragraphs:
            pPr = self._exp._ensure_element(p._p, 'w:pPr', first=True)

            # 零缩进
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                pPr.remove(ind)
            ind = OxmlElement('w:ind')
            for attr in ['left', 'leftChars', 'right', 'rightChars',
                         'firstLine', 'firstLineChars', 'hanging', 'hangingChars']:
                ind.set(qn(f'w:{attr}'), '0')
            pPr.append(ind)

            # 段间距（强制零间距，防止样式默认行距导致表头过高）
            old_sp = pPr.find(qn('w:spacing'))
            if old_sp is not None:
                pPr.remove(old_sp)
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:before'), '0')
            spacing.set(qn('w:after'), '0')
            spacing.set(qn('w:line'), '240')
            spacing.set(qn('w:lineRule'), 'auto')
            pPr.append(spacing)

            # 对齐
            jc = pPr.find(qn('w:jc'))
            if jc is None:
                jc = OxmlElement('w:jc')
                pPr.append(jc)
            jc.set(qn('w:val'), {WD_ALIGN_PARAGRAPH.CENTER: 'center',
                                 WD_ALIGN_PARAGRAPH.RIGHT: 'right',
                                 WD_ALIGN_PARAGRAPH.LEFT: 'left'}.get(align, 'left'))

            for r in p.runs:
                if code_font:
                    r.font.name = 'Consolas'
                elif not r.font.name:
                    r.font.name = normal_font['west']
                if font_size_pt:
                    r.font.size = Pt(font_size_pt)
                elif not r.font.size:
                    r.font.size = normal_font['size']
                if normal_font.get('east'):
                    r.element.rPr.rFonts.set(qn('w:eastAsia'), normal_font['east'])
                if is_header:
                    r.font.bold = True

    def _render_cell_element(self, cell, element, fallback_text: str):
        """Render rich cell contents, promoting line breaks to cell paragraphs."""
        if self._render_direct_blocks(cell, element):
            return

        chunks = self._split_inline_by_breaks(element)
        if not chunks:
            if fallback_text:
                self._write_multiline_text(cell, fallback_text)
            return

        for idx, chunk_html in enumerate(chunks):
            para = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
            if idx == 0:
                para.clear()
            soup = BeautifulSoup(f'<span>{chunk_html}</span>', 'html.parser')
            wrapper = soup.find('span')
            if wrapper is not None:
                self._render_inline_or_image_paragraph(para, wrapper)
            if not para.runs:
                text = BeautifulSoup(chunk_html, 'html.parser').get_text()
                if text:
                    para.add_run(text)

    def _render_direct_blocks(self, cell, element) -> bool:
        """Render block-like children inside an HTML table cell."""
        block_children = self._collect_renderable_blocks(element)
        if not block_children:
            return False

        para_idx = 0
        for child in block_children:
            if child.name == 'p':
                para_idx = self._render_inline_chunks(cell, child, para_idx)
            elif child.name in ('ul', 'ol'):
                ordered = child.name == 'ol'
                counter = 1
                for li in child.find_all('li', recursive=False):
                    para = cell.paragraphs[0] if para_idx == 0 else cell.add_paragraph()
                    if para_idx == 0:
                        para.clear()
                    prefix = f'{counter}. ' if ordered else '● '
                    counter += 1
                    para.add_run(prefix)
                    direct_paragraphs = li.find_all('p', recursive=False)
                    if direct_paragraphs:
                        self._render_inline_or_image_paragraph(para, direct_paragraphs[0])
                        para_idx += 1
                        for extra_p in direct_paragraphs[1:]:
                            extra_para = cell.add_paragraph()
                            self._render_inline_or_image_paragraph(extra_para, extra_p)
                            para_idx += 1
                        continue
                    self._render_inline_or_image_paragraph(para, li)
                    para_idx += 1
            elif child.name == 'pre':
                text = child.get_text()
                for line in text.splitlines() or ['']:
                    para = cell.paragraphs[0] if para_idx == 0 else cell.add_paragraph()
                    if para_idx == 0:
                        para.clear()
                    run = para.add_run(line)
                    run.font.name = 'Consolas'
                    para_idx += 1
        return para_idx > 0

    def _render_inline_chunks(self, cell, element, para_idx: int) -> int:
        chunks = self._split_inline_by_breaks(element)
        if not chunks:
            para = cell.paragraphs[0] if para_idx == 0 else cell.add_paragraph()
            if para_idx == 0:
                para.clear()
            self._render_inline_or_image_paragraph(para, element)
            return para_idx + 1
        for chunk_html in chunks:
            para = cell.paragraphs[0] if para_idx == 0 else cell.add_paragraph()
            if para_idx == 0:
                para.clear()
            soup = BeautifulSoup(f'<span>{chunk_html}</span>', 'html.parser')
            wrapper = soup.find('span')
            if wrapper is not None:
                self._render_inline_or_image_paragraph(para, wrapper)
            para_idx += 1
        return para_idx

    @staticmethod
    def _collect_renderable_blocks(element) -> list:
        """Collect block nodes, including simple wrapper nesting."""
        blocks = [
            child for child in element.children
            if getattr(child, 'name', None) in ('p', 'ul', 'ol', 'pre')
        ]
        if blocks:
            return blocks
        nested_blocks = []
        for block in element.find_all(['p', 'ul', 'ol', 'pre']):
            parent = block.parent
            nested_inside_block = False
            while parent is not None and parent is not element:
                if getattr(parent, 'name', None) in ('p', 'ul', 'ol', 'pre'):
                    nested_inside_block = True
                    break
                parent = parent.parent
            if not nested_inside_block:
                nested_blocks.append(block)
        return nested_blocks

    def _render_inline_or_image_paragraph(self, para, element):
        """Render inline content and real images inside a table paragraph."""
        images = element.find_all('img')
        if not images:
            self._exp._process_inline_runs(para, element)
            return

        for child in element.children:
            if getattr(child, 'name', None) == 'img':
                if not self._add_cell_image(para, child):
                    alt = child.get('alt', '') or child.get('src', '')
                    para.add_run(f'[图片: {alt}]')
            elif isinstance(child, NavigableString):
                if str(child):
                    para.add_run(str(child))
            else:
                soup = BeautifulSoup(f'<span>{child}</span>', 'html.parser')
                wrapper = soup.find('span')
                if wrapper is not None:
                    self._exp._process_inline_runs(para, wrapper)

    def _add_cell_image(self, para, img_tag) -> bool:
        """Insert a local or data URI image into a table cell paragraph."""
        src = img_tag.get('src', '')
        if not src:
            return False

        tmp_path = None
        image_arg = src
        try:
            if src.startswith('data:'):
                header, encoded = src.split(',', 1)
                ext = header.split(';')[0].split('/')[-1] if 'image/' in header else 'png'
                data = base64.b64decode(encoded)
                with tempfile.NamedTemporaryFile(suffix=f'.{ext}', delete=False) as f:
                    f.write(data)
                    tmp_path = f.name
                    image_arg = tmp_path
            elif src.startswith(('http://', 'https://')):
                cached = self._download_remote_image(src)
                if not cached:
                    return False
                image_arg = cached
            elif not Path(src).exists():
                return False

            width = Pt(72)  # about 1 inch; keeps table rows compact
            para.add_run().add_picture(image_arg, width=width)
            return True
        except Exception:
            return False
        finally:
            if tmp_path:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass

    @staticmethod
    def _download_remote_image(src: str) -> str:
        """Download a remote image into a deterministic temp cache."""
        if requests is None:
            return ''
        try:
            cache_dir = Path(tempfile.gettempdir()) / 'techdoc-md-renderer-image-cache'
            cache_dir.mkdir(parents=True, exist_ok=True)
            suffix = Path(src.split('?', 1)[0]).suffix
            if suffix.lower() not in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp'):
                suffix = '.img'
            digest = hashlib.sha256(src.encode('utf-8')).hexdigest()[:16]
            target = cache_dir / f'{digest}{suffix}'
            if target.exists() and target.stat().st_size > 0:
                return str(target)
            resp = requests.get(src, timeout=8)
            resp.raise_for_status()
            content_type = resp.headers.get('content-type', '')
            if not content_type.startswith('image/'):
                return ''
            target.write_bytes(resp.content)
            return str(target)
        except Exception:
            return ''

    @staticmethod
    def _split_inline_by_breaks(element) -> list:
        chunks = []
        current = []
        for child in element.children:
            if getattr(child, 'name', None) == 'br':
                chunk = ''.join(current).strip()
                if chunk:
                    chunks.append(chunk)
                current = []
                continue
            if isinstance(child, NavigableString):
                current.append(str(child))
            else:
                current.append(str(child))
        chunk = ''.join(current).strip()
        if chunk:
            chunks.append(chunk)
        return chunks

    def _write_multiline_text(self, cell, text: str):
        lines = (text or '').splitlines()
        if not lines:
            lines = ['']
        for idx, line in enumerate(lines):
            para = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
            if idx == 0:
                para.clear()
            para.add_run(line)

    def _set_borders(self, table, widths=None):
        """设置表格边框（全边框 0.5pt 灰色）"""
        tbl = table._tbl
        tblPr = self._exp._ensure_element(tbl, 'w:tblPr', first=True)

        tblBorders = tblPr.find(qn('w:tblBorders'))
        if tblBorders is not None:
            tblPr.remove(tblBorders)

        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), str(BORDER_SIZE))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), BORDER_COLOR)
            tblBorders.append(border)
        tblPr.append(tblBorders)

        # 表格宽度（A4 正文区 ~9000 dxa，polish 阶段会智能重分配列宽）
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        total_width = sum(widths) if widths else 9000
        tblW.set(qn('w:w'), str(total_width))
        tblW.set(qn('w:type'), 'dxa')

        # 固定布局，避免 Word 打开后重新自动适配列宽。
        tblLayout = tblPr.find(qn('w:tblLayout'))
        if tblLayout is None:
            tblLayout = OxmlElement('w:tblLayout')
            tblPr.append(tblLayout)
        tblLayout.set(qn('w:type'), 'fixed')

        # 生成 gridCol 元素。若分类器给出宽度，则优先使用场景策略。
        col_count = len(table.columns)
        if col_count > 0:
            for old_grid in tbl.findall(qn('w:tblGrid')):
                tbl.remove(old_grid)
            widths = widths if widths and len(widths) == col_count else None
            tblGrid = OxmlElement('w:tblGrid')
            equal_width = 9000 // col_count
            for ci in range(col_count):
                gc = OxmlElement('w:gridCol')
                gc.set(qn('w:w'), str(widths[ci] if widths else equal_width))
                tblGrid.append(gc)
            # 插入到 tblPr 之后
            tblPr.addnext(tblGrid)

    def _apply_layout(self, table, widths, header_count: int):
        """Synchronize grid widths into cell widths and repeat header rows."""
        if not widths or len(widths) != len(table.columns):
            return

        for row_idx, row in enumerate(table.rows):
            if header_count and row_idx < header_count:
                trPr = self._exp._ensure_element(row._tr, 'w:trPr', first=True)
                tblHeader = trPr.find(qn('w:tblHeader'))
                if tblHeader is None:
                    tblHeader = OxmlElement('w:tblHeader')
                    trPr.append(tblHeader)
                tblHeader.set(qn('w:val'), 'true')

            for col_idx, cell in enumerate(row.cells):
                if col_idx >= len(widths):
                    continue
                tcPr = self._exp._ensure_element(cell._tc, 'w:tcPr', first=True)
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:w'), str(widths[col_idx]))
                tcW.set(qn('w:type'), 'dxa')

    @staticmethod
    def _wrap_long_code_cell(text: str) -> str:
        """Insert visible line breaks into long grammar/code-like table cells."""
        import re

        compact = re.sub(r'\s+', ' ', text.strip())
        if len(compact) <= 120:
            return compact

        # Prefer semantic break points in grammar-ish definitions. These are
        # display breaks only; the text remains readable and copyable.
        parts = re.split(r'(?<=[>}\\]])\s*|(?<=\))\s*|(?<=;)\s*', compact)
        lines = []
        current = ''
        for part in parts:
            if not part:
                continue
            if current and len(current) + len(part) > 48:
                lines.append(current.rstrip())
                current = part
            else:
                current += part
        if current:
            lines.append(current.rstrip())

        if len(lines) <= 1:
            lines = [compact[i:i + 48] for i in range(0, len(compact), 48)]
        return '\n'.join(lines)

    @staticmethod
    def _begin_landscape_section(doc):
        """Start a continuous landscape section for wide technical tables."""
        section = doc.add_section(WD_SECTION_START.CONTINUOUS)
        if section.page_width < section.page_height:
            section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE

    @staticmethod
    def _end_portrait_section(doc):
        """Return following content to portrait orientation."""
        section = doc.add_section(WD_SECTION_START.CONTINUOUS)
        if section.page_width > section.page_height:
            section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.PORTRAIT
