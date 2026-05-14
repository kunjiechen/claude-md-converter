"""Word section, page-break and revision helpers."""

from __future__ import annotations

import copy

from bs4 import BeautifulSoup
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class SectionBuilder:
    """Build document sections that are not regular body content."""

    def __init__(self, exporter):
        self._exp = exporter

    def insert_toc_field(self, doc):
        body = doc.element.body

        title_para = OxmlElement('w:p')
        title_pPr = OxmlElement('w:pPr')
        title_pStyle = OxmlElement('w:pStyle')
        title_pStyle.set(qn('w:val'), 'Heading1')
        title_pPr.append(title_pStyle)
        title_para.append(title_pPr)
        title_run = OxmlElement('w:r')
        title_t = OxmlElement('w:t')
        title_t.set(qn('xml:space'), 'preserve')
        title_t.text = '目录'
        title_run.append(title_t)
        title_para.append(title_run)
        body.insert(0, title_para)

        toc_para = OxmlElement('w:p')
        toc_r = OxmlElement('w:r')
        fldChar_begin = OxmlElement('w:fldChar')
        fldChar_begin.set(qn('w:fldCharType'), 'begin')
        toc_r.append(fldChar_begin)
        toc_para.append(toc_r)

        instr_r = OxmlElement('w:r')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'TOC \\o "1-2" \\h \\z \\u'
        instr_r.append(instr)
        toc_para.append(instr_r)

        sep_r = OxmlElement('w:r')
        fldChar_sep = OxmlElement('w:fldChar')
        fldChar_sep.set(qn('w:fldCharType'), 'separate')
        sep_r.append(fldChar_sep)
        toc_para.append(sep_r)

        end_r = OxmlElement('w:r')
        fldChar_end = OxmlElement('w:fldChar')
        fldChar_end.set(qn('w:fldCharType'), 'end')
        end_r.append(fldChar_end)
        toc_para.append(end_r)
        body.insert(1, toc_para)

        page_break_para = self.page_break_paragraph()
        body.insert(2, page_break_para)

        settings = doc.settings.element
        updateFields = settings.find(qn('w:updateFields'))
        if updateFields is None:
            updateFields = OxmlElement('w:updateFields')
            settings.append(updateFields)
        updateFields.set(qn('w:val'), 'true')

    def add_pagebreak(self, doc):
        para = doc.add_paragraph()
        run = para.add_run()
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        run._r.append(br)

    def add_revision_section(self, doc, revision_html: str):
        if not revision_html:
            return

        soup = BeautifulSoup(revision_html, 'html.parser')
        heading_tag = soup.find('h2')
        table_tag = soup.find('table')
        if not table_tag:
            return

        if heading_tag:
            self._exp._add_heading(heading_tag, doc)
        self._exp._table_builder.build(table_tag, doc)
        self.add_pagebreak(doc)

    @staticmethod
    def page_break_paragraph():
        page_break_p = OxmlElement('w:p')
        pb_r = OxmlElement('w:r')
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        pb_r.append(br)
        page_break_p.append(pb_r)
        return page_break_p

    @staticmethod
    def copy_element(element):
        return copy.deepcopy(element)

