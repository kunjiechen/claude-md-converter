"""Word list builder — stable manual bullet with hanging indent, no tab characters."""

from __future__ import annotations

from bs4 import NavigableString, Tag
from docx.shared import Cm, Pt


class ListBuilder:
    """Render HTML ``ul``/``ol`` lists as Word paragraphs.

    Uses a simple bullet-prefix + hanging-indent approach *without* tab
    characters, because ``\\t`` interacts unpredictably with Word's tab-stop
    engine when content runs use different fonts/metrics than the prefix run.
    """

    BULLET_CHARS = ['●', '◆', '■', '▸']

    def __init__(self, exporter):
        self._exp = exporter

    def build(self, tag: Tag, doc, level: int = 0):
        is_ordered = tag.name == 'ol'
        left = Cm(0.85 + level * 0.65)
        hang = Cm(-0.6)
        counter = 1

        for li in tag.find_all('li', recursive=False):
            list_style = self._choose_style(doc, li)
            para = doc.add_paragraph(style=list_style) if list_style else doc.add_paragraph()
            para.paragraph_format.left_indent = left
            para.paragraph_format.first_line_indent = hang
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)

            task_checked = self._task_checked(li)
            if task_checked is not None:
                prefix = '☑  ' if task_checked else '☐  '
            elif is_ordered:
                prefix = f'{counter}、  '
                counter += 1
            else:
                prefix = self.BULLET_CHARS[level % len(self.BULLET_CHARS)] + '  '

            prefix_run = para.add_run(prefix)
            prefix_run.font.size = Pt(10.5)

            for child in li.children:
                if isinstance(child, NavigableString):
                    text = child.strip()
                    if text:
                        run = para.add_run(text)
                        run.font.size = Pt(10.5)
                elif child.name in ('ul', 'ol', 'input'):
                    continue
                else:
                    self._exp._render_single_inline(para, child)

            if self._is_only_prefix(para):
                body = para._p.getparent()
                if body is not None:
                    body.remove(para._p)

            for nested in li.find_all(['ul', 'ol'], recursive=False):
                self.build(nested, doc, level + 1)

    def _choose_style(self, doc, li: Tag) -> str | None:
        li_children = [c for c in li.children if c.name is not None or str(c).strip()]
        strong_only = (
            len(li_children) == 1
            and getattr(li_children[0], 'name', None) == 'strong'
            and len(li_children[0].get_text(strip=True)) <= 50
        )
        if strong_only and self._exp._has_style(doc, '小标题'):
            return '小标题'
        if self._exp._has_style(doc, '正文2'):
            return '正文2'
        return None

    @staticmethod
    def _task_checked(li: Tag):
        for child in li.children:
            if isinstance(child, NavigableString):
                continue
            if child.name == 'input':
                return child.get('checked')
        return None

    @staticmethod
    def _is_only_prefix(para) -> bool:
        content_runs = [r for r in para.runs if r.text.strip()]
        if len(content_runs) != 1:
            return False
        text = content_runs[0].text.strip()
        bullets = ('●', '◆', '■', '▸', '☑', '☐')
        return text in bullets or text.endswith('、')
