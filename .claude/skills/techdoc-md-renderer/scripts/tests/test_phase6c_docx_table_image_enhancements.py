from pathlib import Path
import base64
import sys
import tempfile
import unittest

from docx import Document as ReadDocx


SCRIPT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPT_DIR))

from parser import MarkdownParser
from core.layout import LayoutPlanner
from core.model import Document, Table, TableCell, TableRow, TextRun
from core.normalize import ast_to_document
from core.rules import PolicyBuilder
from core.semantic import analyze_document
from renderers.base import RenderContext
from renderers.docx import DocxRendererAdapter


ONE_PIXEL_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


class Phase6CDocxTableImageEnhancementsTest(unittest.TestCase):
    def context_from_markdown(self, markdown: str, base_path: Path, *, profile: str = "chip_register_manual") -> RenderContext:
        ast = MarkdownParser().parse(markdown)
        doc = ast_to_document(ast, source_hint=str(base_path / "phase6c.md"))
        analyze_document(doc, base_path=base_path)
        policy = PolicyBuilder().build(doc, document_profile=profile)
        plan = LayoutPlanner().plan(doc, policy)
        return RenderContext(
            document=doc,
            policy=policy,
            layout_plan=plan,
            output_path=base_path / "out.docx",
            diagnostics=doc.metadata.get("semantic_analysis", {}).get("diagnostics", []),
            assets={"base_path": str(base_path)},
        )

    def render_markdown(self, markdown: str, base_path: Path, *, profile: str = "chip_register_manual"):
        context = self.context_from_markdown(markdown, base_path, profile=profile)
        result = DocxRendererAdapter().render(context)
        self.assertTrue(result.success, result.error)
        return result, ReadDocx(str(context.output_path))

    def test_table_kind_styles_and_width_intent_for_register_bitfield_interface_generic(self):
        markdown = (
            "<!-- table: register -->\n"
            "| Address | Register | Bits | Access | Reset | Description |\n"
            "| --- | --- | --- | --- | --- | --- |\n"
            "| 0x00 | CTRL | [3:0] | RW | 0x0 | Control register |\n\n"
            "<!-- table: bitfield -->\n"
            "| Bits | Field | Access | Reset | Description |\n"
            "| --- | --- | --- | --- | --- |\n"
            "| [7:4] | MODE | RW | 0 | 工作模式 |\n\n"
            "<!-- table: interface -->\n"
            "| 元素名称 | 基本结构 | 描述 |\n"
            "| --- | --- | --- |\n"
            "| RequestId | `<Id>` | 请求标识 |\n\n"
            "### Generic\n\n"
            "| A | B | C | D | E |\n"
            "| --- | --- | --- | --- | --- |\n"
            "| one | two | three | four | five |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render_markdown(markdown, Path(tmp))
        self.assertEqual(len(docx.tables), 4)
        xml = "\n".join(table._tbl.xml for table in docx.tables)
        self.assertIn("D9EAF7", xml)
        self.assertIn("E2F0D9", xml)
        self.assertIn("EDE7F6", xml)
        self.assertIn("w:tblHeader", xml)
        self.assertIn("w:tcW", xml)
        self.assertTrue(any(d.get("code") == "docx_table_wrap_policy_recorded" for d in result.diagnostics))

    def test_wide_table_overflow_and_landscape_recommendation_diagnostics(self):
        markdown = (
            "<!-- table: register -->\n"
            "| Address | Register | Bits | Access | Reset | Clock | Domain | Description |\n"
            "| --- | --- | --- | --- | --- | --- | --- | --- |\n"
            "| 0x00 | CTRL | [3:0] | RW | 0x0 | APB | SYS | Very long description for layout risk. |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.render_markdown(markdown, Path(tmp))
        codes = {d.get("code") for d in result.diagnostics}
        self.assertIn("docx_table_layout_review", codes)
        self.assertIn("docx_landscape_recommendation_not_applied", codes)

    def test_colspan_and_rowspan_emit_safe_fallback_diagnostics(self):
        doc = Document(blocks=[
            Table(rows=[
                TableRow(header=True, cells=[TableCell(text="H1", children=[TextRun(text="H1")], colspan=2), TableCell(text="H2", children=[TextRun(text="H2")])]),
                TableRow(cells=[TableCell(text="A", children=[TextRun(text="A")], rowspan=2), TableCell(text="B", children=[TextRun(text="B")])]),
            ])
        ])
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            analyze_document(doc, base_path=base)
            policy = PolicyBuilder().build(doc, document_profile="lightweight_tech_note")
            plan = LayoutPlanner().plan(doc, policy)
            result = DocxRendererAdapter().render(RenderContext(document=doc, policy=policy, layout_plan=plan, output_path=base / "span.docx"))
        self.assertTrue(result.success)
        self.assertTrue(any(d.get("code") == "docx_complex_table_span_unsupported" for d in result.diagnostics))

    def test_local_png_missing_svg_caption_and_mixed_inline_image(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            (base / "tiny.png").write_bytes(base64.b64decode(ONE_PIXEL_PNG))
            (base / "icon.svg").write_text("<svg></svg>", encoding="utf-8")
            markdown = (
                "请参考 ![内联图](tiny.png) 完成连接。\n\n"
                "![图 1 本地图片](tiny.png)\n\n"
                "![缺失](missing.png)\n\n"
                "![矢量图](icon.svg)\n"
            )
            result, docx = self.render_markdown(markdown, base)
        text = "\n".join(p.text for p in docx.paragraphs)
        self.assertIn("请参考", text)
        self.assertIn("完成连接", text)
        self.assertIn("图 1 本地图片", text)
        self.assertIn("Image unavailable: missing.png", text)
        self.assertIn("SVG image unsupported", text)
        self.assertGreaterEqual(len(docx.inline_shapes), 2)
        codes = {d.get("code") for d in result.diagnostics}
        self.assertIn("docx_image_scaling_applied", codes)
        self.assertIn("docx_image_placeholder", codes)
        self.assertIn("docx_svg_unsupported", codes)

    def test_existing_bare_image_paths_render_and_missing_paths_remain_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            (base / "images").mkdir()
            (base / "images" / "tiny.png").write_bytes(base64.b64decode(ONE_PIXEL_PNG))
            markdown = (
                "images/tiny.png\n\n"
                "images/missing.png\n\n"
                "| 图 | 说明 |\n"
                "| --- | --- |\n"
                "| images/tiny.png | symbol |\n"
            )
            result, docx = self.render_markdown(markdown, base)
        text = "\n".join(p.text for p in docx.paragraphs)
        table_text = "\n".join(cell.text for table in docx.tables for row in table.rows for cell in row.cells)
        self.assertNotIn("images/tiny.png", text + table_text)
        self.assertIn("images/missing.png", text)
        self.assertGreaterEqual(len(docx.inline_shapes), 2)
        self.assertTrue(any(d.get("code") == "docx_image_scaling_applied" for d in result.diagnostics))

    def test_image_only_table_cells_use_consistent_centered_width(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            (base / "tiny.png").write_bytes(base64.b64decode(ONE_PIXEL_PNG))
            markdown = (
                "<table><tr><th>示例</th><th>示例</th></tr>"
                "<tr>"
                "<td><img src=\"tiny.png\" style=\"width:3in\" /></td>"
                "<td><img src=\"tiny.png\" style=\"width:.5in\" /></td>"
                "</tr></table>"
            )
            _, docx = self.render_markdown(markdown, base)
        widths = [round(shape.width / 914400, 2) for shape in docx.inline_shapes]
        self.assertGreaterEqual(len(widths), 2)
        self.assertEqual(widths[-2:], [widths[-2], widths[-2]])
        for table in docx.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.paragraphs and not cell.text.strip():
                        self.assertEqual(cell.paragraphs[0].alignment, 1)


if __name__ == "__main__":
    unittest.main()
