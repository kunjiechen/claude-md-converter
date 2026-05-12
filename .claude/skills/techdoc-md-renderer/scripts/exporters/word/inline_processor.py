"""HTML 内联元素 → Word runs 处理器

将 BeautifulSoup4 解析的 HTML 内联元素（strong/em/code/a/kbd/sub/sup/mark/span/input）
渲染为 python-docx 的格式化 Run 对象。
"""

from bs4 import Tag, NavigableString
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class InlineProcessor:
    """将 HTML 内联元素转换为 Word runs"""

    # 格式属性映射
    _FORMAT_ATTR = {'strong': 'bold', 'em': 'italic', 'del': 'strike', 'ins': 'underline'}

    def __init__(self, exporter):
        """
        Args:
            exporter: WordExporter 实例，提供 _add_hyperlink / _add_footnote_reference / _ensure_element
        """
        self._exp = exporter

    def process(self, para, tag: Tag):
        """处理 HTML 元素的内联子元素，为 para 添加 runs"""
        self._render_children(para, tag)

    def _render_children(self, para, element):
        """递归渲染内联子元素"""
        for child in element.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text.strip() or text == ' ':
                    para.add_run(text)
            else:
                self._render_single(para, child)

    def _render_single(self, para, element):
        """渲染单个内联元素，支持嵌套格式叠加"""
        name = element.name
        if name is None:
            return

        if name in ('strong', 'em', 'del', 'ins'):
            self._render_nested(para, element, name)
        elif name == 'code':
            run = para.add_run(element.get_text())
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            shd = self._exp._ensure_element(run._r, 'w:rPr', first=True, use_inner=True)
            shd_el = OxmlElement('w:shd')
            shd_el.set(qn('w:val'), 'clear')
            shd_el.set(qn('w:color'), 'auto')
            shd_el.set(qn('w:fill'), 'F0F0F0')
            shd.insert(0, shd_el)
        elif name == 'a':
            text = element.get_text()
            href = element.get('href', '')
            if href:
                self._exp._add_hyperlink(para, text, href)
            else:
                para.add_run(text)
        elif name == 'kbd':
            run = para.add_run(element.get_text())
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif name == 'sub':
            run = para.add_run(element.get_text())
            run.font.subscript = True
        elif name == 'sup':
            classes = element.get('class', [])
            if 'footnote-ref' in classes:
                label = element.get_text().strip('[] \n')
                self._exp._add_footnote_reference(para, label)
            else:
                run = para.add_run(element.get_text())
                run.font.superscript = True
        elif name == 'mark':
            run = para.add_run(element.get_text())
            shd = self._exp._ensure_element(run._r, 'w:rPr', first=True, use_inner=True)
            shd_el = OxmlElement('w:shd')
            shd_el.set(qn('w:val'), 'clear')
            shd_el.set(qn('w:color'), 'auto')
            shd_el.set(qn('w:fill'), 'FFFF00')
            shd.insert(0, shd_el)
        elif name == 'br':
            para.add_run('\n')
        elif name == 'img':
            alt = element.get('alt', '[图片]')
            para.add_run(f'[图片: {alt}]')
        elif name == 'span':
            classes = element.get('class', [])
            if 'math' in classes:
                text = element.get_text().strip()
                if text.startswith('\\('):
                    text = text[2:]
                if text.endswith('\\)'):
                    text = text[:-2]
                run = para.add_run(text.strip())
                run.font.name = 'Cambria Math'
                run.font.italic = True
                run.font.size = Pt(11)
            else:
                self._render_children(para, element)
        elif name == 'input':
            checked = element.get('checked')
            para.add_run('☑ ' if checked is not None else '☐ ')
        else:
            t = element.get_text()
            if t:
                para.add_run(t)

    def _render_nested(self, para, element, fmt_name):
        """递归渲染 strong/em/del/ins，支持嵌套子元素叠加格式"""
        attr = self._FORMAT_ATTR[fmt_name]
        for child in element.children:
            if isinstance(child, NavigableString):
                t = str(child)
                if t.strip() or t == ' ':
                    run = para.add_run(t)
                    setattr(run.font, attr, True)
            elif child.name in self._FORMAT_ATTR:
                self._render_nested(para, child, child.name)
                if para.runs:
                    setattr(para.runs[-1].font, attr, True)
            else:
                self._render_single(para, child)
                if para.runs:
                    setattr(para.runs[-1].font, attr, True)
