from pathlib import Path
import base64
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document as ReadDocx


SCRIPT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPT_DIR))

from api import Converter
from parser import MarkdownParser
from core.layout import LayoutPlanner
from core.model import Document, UnsupportedBlock
from core.normalize import ast_to_document
from core.rules import PolicyBuilder
from core.semantic import analyze_document
from renderers.base import RenderContext
from renderers.docx import DocxRendererAdapter


ONE_PIXEL_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


class Phase6BDocxAdapterTest(unittest.TestCase):
    def build_context(self, markdown: str, output_path: Path, *, base_path: Path, document_profile: str = "chip_register_manual"):
        ast = MarkdownParser().parse(markdown)
        doc = ast_to_document(ast, source_hint="docx-adapter-test.md")
        analyze_document(doc, base_path=base_path)
        policy = PolicyBuilder().build(doc, document_profile=document_profile)
        plan = LayoutPlanner().plan(doc, policy)
        return RenderContext(
            document=doc,
            policy=policy,
            layout_plan=plan,
            output_path=output_path,
            diagnostics=doc.metadata.get("semantic_analysis", {}).get("diagnostics", []),
            assets={"base_path": str(base_path)},
            options={"doc_title": "DOCX Adapter Test"},
        )

    def render_docx(self, markdown: str, *, base_path: Path, document_profile: str = "chip_register_manual"):
        out = base_path / "adapter.docx"
        context = self.build_context(markdown, out, base_path=base_path, document_profile=document_profile)
        result = DocxRendererAdapter().render(context)
        self.assertTrue(result.success, result.error)
        self.assertTrue(out.exists())
        return result, ReadDocx(str(out))

    @staticmethod
    def doc_text(docx) -> str:
        return "\n".join(p.text for p in docx.paragraphs)

    def test_simple_heading_and_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render_docx("# 标题\n\n正文段落。", base_path=Path(tmp), document_profile="lightweight_tech_note")
        self.assertIn("标题", self.doc_text(docx))
        self.assertIn("正文段落", self.doc_text(docx))
        self.assertEqual(result.fidelity_level, "review")
        self.assertTrue(any(d.get("code") == "unsupported_section_layout" for d in result.diagnostics))

    def test_inline_bold_italic_code_link(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, docx = self.render_docx("这是 **粗体**、*斜体*、`CODE` 和 [链接](https://example.com)。", base_path=Path(tmp))
        text = self.doc_text(docx)
        self.assertIn("粗体", text)
        self.assertIn("斜体", text)
        self.assertIn("CODE", text)
        self.assertIn("https://example.com", text)

    def test_mixed_text_image_paragraph_preserves_text_and_image_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render_docx("请参考 ![缺失图](missing.png) 完成连接。", base_path=Path(tmp))
        text = self.doc_text(docx)
        self.assertIn("请参考", text)
        self.assertIn("Image unavailable: missing.png", text)
        self.assertIn("完成连接", text)
        self.assertTrue(any(d.get("code") == "docx_image_placeholder" for d in result.diagnostics))

    def test_local_image_can_be_inserted(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            (base / "tiny.png").write_bytes(base64.b64decode(ONE_PIXEL_PNG))
            result, docx = self.render_docx("![tiny](tiny.png)", base_path=base)
        self.assertFalse(any(d.get("code") == "docx_image_placeholder" for d in result.diagnostics))
        self.assertGreaterEqual(len(docx.inline_shapes), 1)

    def test_code_block_identity_is_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, docx = self.render_docx("```c\nuint8_t value = read_reg(0x00);\n```", base_path=Path(tmp))
        self.assertIn("uint8_t value", self.doc_text(docx))

    def test_simple_and_register_tables_render_without_reclassifying(self):
        markdown = (
            "| 名称 | 说明 |\n| --- | --- |\n| ENABLE | 使能 |\n\n"
            "<!-- table: register -->\n"
            "| Address | Register | Bits | Access | Reset | Description |\n"
            "| --- | --- | --- | --- | --- | --- |\n"
            "| 0x00 | CTRL | [3:0] | RW | 0x0 | Control register |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            context = self.build_context(markdown, Path(tmp) / "tables.docx", base_path=Path(tmp))
            with patch("analyzers.table_classifier.TableClassifier.classify", side_effect=AssertionError("classification during docx render")):
                result = DocxRendererAdapter().render(context)
            docx = ReadDocx(str(context.output_path))
        self.assertTrue(result.success)
        self.assertEqual(len(docx.tables), 2)
        self.assertIn("CTRL", docx.tables[1].cell(1, 1).text)

    def test_wide_table_layout_intent_produces_review_diagnostic(self):
        markdown = (
            "<!-- table: register -->\n"
            "| Address | Register | Bits | Access | Reset | Clock | Domain | Description |\n"
            "| --- | --- | --- | --- | --- | --- | --- | --- |\n"
            "| 0x00 | CTRL | [3:0] | RW | 0x0 | APB | SYS | Long description for layout risk. |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.render_docx(markdown, base_path=Path(tmp))
        self.assertTrue(any(d.get("code") == "docx_table_layout_review" for d in result.diagnostics))
        self.assertEqual(result.fidelity_level, "review")

    def test_missing_image_raw_html_pagebreak_and_unsupported_diagnostics(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render_docx(
                "![缺失](missing.png)\n\n"
                "<div>raw html</div>\n\n"
                "<!-- pagebreak -->\n",
                base_path=Path(tmp),
            )
        text = self.doc_text(docx)
        self.assertIn("Image unavailable: missing.png", text)
        self.assertIn("raw html", text)
        codes = {d.get("code") for d in result.diagnostics}
        self.assertIn("docx_image_placeholder", codes)
        self.assertIn("docx_raw_html_fallback", codes)

    def test_list_and_blockquote_basic_support(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, docx = self.render_docx("- 一级\n  - 二级\n\n> 引用段落", base_path=Path(tmp))
        text = self.doc_text(docx)
        self.assertIn("一级", text)
        self.assertIn("二级", text)
        self.assertIn("引用段落", text)

    def test_unsupported_node_diagnostic(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document(blocks=[UnsupportedBlock(original_type="custom", content="payload")])
            analyze_document(doc, base_path=Path(tmp))
            policy = PolicyBuilder().build(doc, document_profile="lightweight_tech_note")
            plan = LayoutPlanner().plan(doc, policy)
            result = DocxRendererAdapter().render(RenderContext(document=doc, policy=policy, layout_plan=plan, output_path=Path(tmp) / "unsupported.docx"))
        self.assertTrue(result.success)
        self.assertTrue(any(d.get("code") == "docx_unsupported_block" for d in result.diagnostics))

    def test_word_v2_conversion_outputs_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "sample.md"
            md.write_text("# 标题\n\n正文。", encoding="utf-8")
            result = Converter().convert_file(md, format="word", output_path=Path(tmp) / "v2.docx", pipeline="v2")
        self.assertTrue(result.success, result.error)


if __name__ == "__main__":
    unittest.main()
