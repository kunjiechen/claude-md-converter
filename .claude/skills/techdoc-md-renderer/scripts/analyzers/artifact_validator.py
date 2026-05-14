"""Final artifact validation for HTML, Word, and PDF outputs."""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
import json
import re
from typing import Any, Dict, List

try:
    from .visual_validator import VisualValidator
except ImportError:  # pragma: no cover - direct script import fallback
    from visual_validator import VisualValidator  # type: ignore


@dataclass
class ArtifactIssue:
    severity: str
    category: str
    message: str
    location: str = ""


@dataclass
class ArtifactValidationReport:
    file_path: str
    format: str
    ok: bool = True
    metrics: Dict[str, Any] = field(default_factory=dict)
    issues: List[ArtifactIssue] = field(default_factory=list)

    @property
    def critical_count(self) -> int:
        return sum(1 for i in self.issues if i.severity == "critical")

    @property
    def warning_count(self) -> int:
        return sum(1 for i in self.issues if i.severity != "critical")

    def to_dict(self) -> Dict[str, Any]:
        payload = asdict(self)
        payload["critical_count"] = self.critical_count
        payload["warning_count"] = self.warning_count
        return payload


class ArtifactValidator:
    """Read-only product-level validation for generated artifacts."""

    def validate(self, file_path: str) -> ArtifactValidationReport:
        p = Path(file_path)
        if not p.exists():
            report = ArtifactValidationReport(str(p), "unknown", ok=False)
            self._add(report, "critical", "file", f"输出文件不存在: {file_path}")
            return report

        suffix = p.suffix.lower()
        if suffix in (".html", ".htm"):
            report = self._validate_html(p)
        elif suffix == ".docx":
            report = self._validate_docx(p)
        elif suffix == ".pdf":
            report = self._validate_pdf(p)
        else:
            report = ArtifactValidationReport(str(p), suffix.lstrip(".") or "unknown", ok=False)
            self._add(report, "critical", "file", f"不支持的产物格式: {suffix}")

        report.ok = report.critical_count == 0
        return report

    def _validate_html(self, file_path: Path) -> ArtifactValidationReport:
        report = ArtifactValidationReport(str(file_path), "html")
        html = file_path.read_text(encoding="utf-8", errors="ignore")
        report.metrics["size_bytes"] = file_path.stat().st_size
        report.metrics["tables"] = len(re.findall(r"<table\b", html, re.I))
        report.metrics["images"] = len(re.findall(r"<img\b", html, re.I))
        report.metrics["headings"] = len(re.findall(r"<h[1-6]\b", html, re.I))

        if file_path.stat().st_size < 512:
            self._add(report, "critical", "file", "HTML 文件过小，可能生成失败")
        if "<body" not in html.lower():
            self._add(report, "critical", "structure", "HTML 缺少 body")
        body_text = re.sub(r"<[^>]+>", " ", html)
        if len(body_text.strip()) < 20:
            self._add(report, "critical", "content", "HTML 正文内容过少")
        if "<style" not in html.lower():
            self._add(report, "warning", "style", "HTML 未包含内联样式，阅读效果可能不完整")
        if '<pre class="mermaid"' in html and "mermaid" not in html.lower():
            self._add(report, "critical", "mermaid", "HTML 含 Mermaid 代码但未配置渲染脚本")
        if re.search(r"<script\b[^>]*>", html, re.I) and "mermaid" not in html.lower():
            self._add(report, "warning", "security", "HTML 含脚本，请确认输入源可信")
        return report

    def _validate_docx(self, file_path: Path) -> ArtifactValidationReport:
        report = ArtifactValidationReport(str(file_path), "docx")
        report.metrics["size_bytes"] = file_path.stat().st_size
        if file_path.stat().st_size < 2048:
            self._add(report, "critical", "file", "DOCX 文件过小，可能生成失败")

        try:
            from docx import Document
        except ImportError:
            self._add(report, "warning", "dependency", "python-docx 不可用，跳过 DOCX 深度校验")
            return report

        try:
            doc = Document(str(file_path))
        except Exception as exc:
            self._add(report, "critical", "file", f"DOCX 无法打开: {exc}")
            return report

        paragraph_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        table_count = len(doc.tables)
        rel_images = [
            rel for rel in doc.part.rels.values()
            if "image" in getattr(rel, "reltype", "")
        ]
        report.metrics.update({
            "paragraphs": len(paragraph_texts),
            "tables": table_count,
            "sections": len(doc.sections),
            "images": len(rel_images),
        })

        if len(paragraph_texts) + table_count == 0:
            self._add(report, "critical", "content", "DOCX 正文和表格均为空")
        if paragraph_texts and paragraph_texts[0] == "目录" and len(paragraph_texts) <= 2 and table_count == 0:
            self._add(report, "warning", "toc", "DOCX 可能只有未更新目录，没有正文内容")
        for idx, table in enumerate(doc.tables, 1):
            if len(table.rows) == 0 or len(table.columns) == 0:
                self._add(report, "critical", "table", "DOCX 存在空表格", f"表格 {idx}")
            if len(table.columns) >= 8:
                self._add(report, "warning", "table", f"DOCX 表格列数较多: {len(table.columns)}", f"表格 {idx}")
        self._attach_visual_validation(report, file_path)
        return report

    def _validate_pdf(self, file_path: Path) -> ArtifactValidationReport:
        report = ArtifactValidationReport(str(file_path), "pdf")
        size = file_path.stat().st_size
        report.metrics["size_bytes"] = size
        if size < 1024:
            self._add(report, "critical", "file", "PDF 文件过小，可能生成失败")

        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from PyPDF2 import PdfReader  # type: ignore
            except ImportError:
                self._add(report, "warning", "dependency", "未安装 pypdf，跳过 PDF 页数/文本校验")
                return report

        try:
            reader = PdfReader(str(file_path))
            page_count = len(reader.pages)
            report.metrics["pages"] = page_count
            if page_count == 0:
                self._add(report, "critical", "content", "PDF 没有页面")
                return report
            sample_text = ""
            for page in reader.pages[: min(3, page_count)]:
                try:
                    sample_text += page.extract_text() or ""
                except Exception:
                    continue
            report.metrics["sample_text_length"] = len(sample_text.strip())
            if len(sample_text.strip()) < 20:
                self._add(report, "warning", "content", "PDF 前几页可提取文本过少，需人工确认是否为空白或扫描图")
        except Exception as exc:
            self._add(report, "critical", "file", f"PDF 无法打开: {exc}")
        self._attach_pdf_backend_info(report, file_path)
        if report.critical_count == 0:
            self._attach_visual_validation(report, file_path)
        return report

    def _attach_pdf_backend_info(self, report: ArtifactValidationReport, file_path: Path):
        sidecar = file_path.with_suffix(file_path.suffix + ".backend.json")
        if not sidecar.exists():
            return
        try:
            payload = json.loads(sidecar.read_text(encoding="utf-8"))
        except Exception:
            self._add(report, "warning", "pdf_backend", "PDF 后端信息文件无法读取")
            return
        report.metrics["pdf_backend"] = payload
        backend = payload.get("backend", "unknown")
        warnings = payload.get("warnings") or []
        if payload.get("degraded"):
            message = f"PDF 使用降级后端生成: {backend}"
            if warnings:
                message += f"；{warnings[-1]}"
            self._add(report, "warning", "pdf_backend", message)

    def _attach_visual_validation(self, report: ArtifactValidationReport, file_path: Path):
        visual = VisualValidator().validate(str(file_path))
        report.metrics["visual_validation"] = visual.to_dict()
        for issue in visual.issues:
            self._add(
                report,
                issue.severity,
                f"visual:{issue.category}",
                issue.message,
                issue.location,
            )

    @staticmethod
    def _add(report: ArtifactValidationReport, severity: str, category: str,
             message: str, location: str = ""):
        report.issues.append(ArtifactIssue(severity, category, message, location))
