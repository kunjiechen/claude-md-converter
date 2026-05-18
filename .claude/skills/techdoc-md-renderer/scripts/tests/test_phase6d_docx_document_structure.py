from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document as ReadDocx


SCRIPT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPT_DIR))

from api import Converter
from parser import MarkdownParser
from core.layout import LayoutPlanner
from core.normalize import ast_to_document
from core.rules import PolicyBuilder
from core.semantic import analyze_document
from renderers.base import RenderContext
from renderers.docx import DocxRendererAdapter


class Phase6DDocxDocumentStructureTest(unittest.TestCase):
    def build_context(self, markdown: str, base_path: Path, *, profile: str = "automotive_formal_spec", mutate_policy=None) -> RenderContext:
        ast = MarkdownParser().parse(markdown)
        doc = ast_to_document(ast, source_hint="phase6d.md")
        analyze_document(doc, base_path=base_path)
        policy = PolicyBuilder().build(doc, document_profile=profile)
        if mutate_policy:
            mutate_policy(policy)
        plan = LayoutPlanner().plan(doc, policy)
        return RenderContext(
            document=doc,
            policy=policy,
            layout_plan=plan,
            output_path=base_path / "out.docx",
            diagnostics=doc.metadata.get("semantic_analysis", {}).get("diagnostics", []),
            options={"doc_title": "结构测试"},
        )

    def render(self, markdown: str, base_path: Path, **kwargs):
        context = self.build_context(markdown, base_path, **kwargs)
        result = DocxRendererAdapter().render(context)
        self.assertTrue(result.success, result.error)
        return result, ReadDocx(str(context.output_path))

    @staticmethod
    def text(docx) -> str:
        return "\n".join(p.text for p in docx.paragraphs)

    def test_heading_numbering_h1_to_h4(self):
        markdown = "# A\n\n## B\n\n### C\n\n#### D\n"
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render(markdown, Path(tmp))
        text = self.text(docx)
        self.assertIn("1 A", text)
        self.assertIn("1.1 B", text)
        self.assertIn("1.1.1 C", text)
        self.assertIn("1.1.1.1 D", text)
        self.assertTrue(any(d.get("code") == "heading_numbering_applied" for d in result.diagnostics))

    def test_source_numbered_headings_can_be_stripped_by_policy(self):
        def mutate(policy):
            policy.document_policy["source_number_stripping"] = "strip"

        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render("# 1.2 Source Number\n", Path(tmp), mutate_policy=mutate)
        text = self.text(docx)
        self.assertIn("1 Source Number", text)
        self.assertNotIn("1 1.2 Source Number", text)
        self.assertTrue(any(d.get("code") == "source_number_stripping_applied" for d in result.diagnostics))

    def test_toc_enabled_generates_visible_entries(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render("# 正文\n\n## 子节\n", Path(tmp))
        text = self.text(docx)
        self.assertIn("目录", text)
        self.assertIn("1 正文", text)
        self.assertIn("1.1 子节", text)
        self.assertTrue(any(d.get("code") == "docx_static_toc_generated" for d in result.diagnostics))

    def test_revision_table_exists_is_rendered_without_missing_required_diagnostic(self):
        markdown = (
            "# Spec\n\n"
            "<!-- table: revision -->\n"
            "| 版次 | 修订内容 | 修订日期 |\n"
            "| --- | --- | --- |\n"
            "| A/0 | 初版 | 2026-01-01 |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render(markdown, Path(tmp))
        self.assertEqual(len(docx.tables), 1)
        self.assertFalse(any(d.get("code") == "missing_required_revision_history" and d.get("severity") == "warning" for d in result.diagnostics))

    def test_revision_required_but_missing_reports_review_diagnostic(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.render("# Spec\n\n正文。", Path(tmp))
        self.assertTrue(any(d.get("code") == "missing_required_revision_history" for d in result.diagnostics))

    def test_lightweight_profile_does_not_require_revision(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.render("# Note\n\n正文。", Path(tmp), profile="lightweight_tech_note")
        self.assertFalse(any(d.get("code") == "missing_required_revision_history" for d in result.diagnostics))

    def test_pagebreak_footer_page_number_and_a4_margin_profile(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, docx = self.render("# Spec\n\n正文\n\n<!-- pagebreak -->\n\n下一页", Path(tmp))
        section = docx.sections[0]
        self.assertEqual(round(section.page_width.cm, 1), 21.0)
        self.assertEqual(round(section.page_height.cm, 1), 29.7)
        self.assertAlmostEqual(section.top_margin.cm, 2.0, places=1)
        self.assertAlmostEqual(section.bottom_margin.cm, 1.5, places=1)
        self.assertAlmostEqual(section.left_margin.cm, 2.0, places=1)
        self.assertAlmostEqual(section.right_margin.cm, 2.0, places=1)
        footer_xml = section.footer._element.xml
        self.assertIn("PAGE", footer_xml)
        body_xml = docx._element.body.xml
        self.assertIn('w:type="page"', body_xml)
        self.assertTrue(any(d.get("code") == "footer_page_number_applied" for d in result.diagnostics))

    def test_header_title_from_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            _, docx = self.render("# Spec\n\n正文。", Path(tmp))
        self.assertIn("结构测试", docx.sections[0].header.paragraphs[0].text)

    def test_landscape_recommendation_diagnostic(self):
        markdown = (
            "# Spec\n\n"
            "<!-- table: register -->\n"
            "| Address | Register | Bits | Access | Reset | Clock | Domain | Description |\n"
            "| --- | --- | --- | --- | --- | --- | --- | --- |\n"
            "| 0x00 | CTRL | [3:0] | RW | 0x0 | APB | SYS | Long description. |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.render(markdown, Path(tmp), profile="chip_register_manual")
        self.assertTrue(any(d.get("code") in ("unsupported_section_layout", "docx_landscape_recommendation_not_applied") for d in result.diagnostics))

    def test_docx_v2_conversion_outputs_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "sample.md"
            md.write_text("# 标题\n\n正文。", encoding="utf-8")
            result = Converter().convert_file(md, format="word", output_path=Path(tmp) / "v2.docx", pipeline="v2")
        self.assertTrue(result.success, result.error)


if __name__ == "__main__":
    unittest.main()
