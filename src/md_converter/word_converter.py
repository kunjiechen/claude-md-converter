"""
Word文档转换器
将Markdown AST转换为Word文档
"""

from typing import Any, Dict, List
from pathlib import Path
from urllib.request import urlopen
from urllib.error import URLError
import tempfile
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .converter import BaseConverter
from .parser import MarkdownParser
from .flowchart_renderer import FlowchartProcessor


class WordConverter(BaseConverter):
    """Word文档转换器"""

    def __init__(self, **options):
        """
        初始化Word转换器

        Args:
            **options: 转换选项
                - template: Word模板文件路径(.dotx)
                - output_dir: 输出目录
                - font: 默认字体
                - font_size: 默认字号
                - flowchart_enabled: 是否启用流程图
        """
        super().__init__(**options)
        self.font = options.get('font', '宋体')
        self.font_size = options.get('font_size', 12)
        self.doc = None
        self.input_dir = None  # 输入文件所在目录，用于解析相对路径
        self._footnote_counter = 0  # 脚注计数器

        # 流程图处理器
        self.flowchart_enabled = options.get('flowchart_enabled', True)
        self.flowchart_processor = FlowchartProcessor(**options) if self.flowchart_enabled else None
        self.flowchart_counter = 0

    def convert(self, ast: List[Dict[str, Any]], output_path: str) -> bool:
        """
        将AST转换为Word文档

        Args:
            ast: Markdown AST节点列表
            output_path: 输出Word文件路径

        Returns:
            转换是否成功
        """
        try:
            # 创建文档
            self.doc = self._create_document()

            # 遍历AST节点
            for node in ast:
                self._process_node(node)

            # 保存文档
            self.doc.save(output_path)
            self.log(f"Word文档已生成: {output_path}")
            return True

        except Exception as e:
            self.log(f"转换失败: {e}", "error")
            return False

    def convert_file(self, input_path: str, output_path: str) -> bool:
        """
        转换Markdown文件为Word文档

        Args:
            input_path: 输入Markdown文件路径
            output_path: 输出Word文件路径

        Returns:
            转换是否成功
        """
        try:
            # 验证输入文件
            self.validate_input(input_path)

            # 记录输入文件目录，用于解析相对路径
            self.input_dir = Path(input_path).parent

            # 解析Markdown
            parser = MarkdownParser()
            ast = parser.parse_file(input_path)

            # 转换为Word
            return self.convert(ast, output_path)

        except Exception as e:
            self.log(f"转换失败: {e}", "error")
            return False

    def _create_document(self):
        """
        创建Word文档

        Returns:
            Word文档对象
        """
        if self.template_path:
            # 加载模板
            doc = Document(self.template_path)
        else:
            # 创建新文档
            doc = Document()

            # 设置默认字体
            style = doc.styles['Normal']
            font = style.font
            font.name = self.font
            font.size = Pt(self.font_size)

            # 设置中文字体
            style.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

        return doc

    def _process_node(self, node: Dict[str, Any]):
        """
        处理AST节点

        Args:
            node: AST节点
        """
        node_type = node.get('type')

        if node_type == 'heading':
            self._add_heading(node)
        elif node_type == 'paragraph':
            self._add_paragraph(node)
        elif node_type == 'list':
            self._add_list(node)
        elif node_type == 'table':
            self._add_table(node)
        elif node_type == 'code_block':
            # 检查是否为流程图
            if self.flowchart_enabled and self.flowchart_processor:
                content = node.get('content', '')
                chart_type = self.flowchart_processor.detect_flowchart(content)
                if chart_type:
                    self._add_flowchart(node, chart_type)
                else:
                    self._add_code_block(node)
            else:
                self._add_code_block(node)
        elif node_type == 'blockquote':
            self._add_blockquote(node)
        elif node_type == 'image':
            self._add_image(node)
        elif node_type == 'hr':
            self._add_hr(node)
        elif node_type == 'footnote_block':
            self._add_footnote_block(node)

    def _add_heading(self, node: Dict[str, Any]):
        """
        添加标题

        Args:
            node: 标题节点
        """
        level = node.get('level', 1)
        content = node.get('content', '')
        segments = node.get('children', [])

        # 添加标题
        heading = self.doc.add_heading('', level=level)

        if segments:
            # 有内联格式，逐段添加
            self._process_inline_content(heading, content, segments)
        else:
            # 纯文本标题
            run = heading.add_run(content)

        # 设置中文字体
        for run in heading.runs:
            run.font.name = self.font
            run.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

    def _add_paragraph(self, node: Dict[str, Any]):
        """
        添加段落

        Args:
            node: 段落节点
        """
        content = node.get('content', '')
        segments = node.get('children', [])

        # 添加段落
        para = self.doc.add_paragraph()

        # 处理内联元素
        self._process_inline_content(para, content, segments if segments else None)

        # 设置中文字体（仅对非格式化的run补设）
        for run in para.runs:
            if not run.font.name:
                run.font.name = self.font
                run.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

    def _process_inline_content(self, para, content: str, segments: List[Dict[str, Any]] = None):
        """
        处理内联内容

        Args:
            para: 段落对象
            content: 纯文本内容（当segments为空时使用）
            segments: 结构化内联段列表（来自parser的_parse_inline_segments）
        """
        if not segments:
            run = para.add_run(content)
            run.font.name = self.font
            run.font.size = Pt(self.font_size)
            return

        for seg in segments:
            seg_type = seg.get("type", "text")

            if seg_type == "text":
                text = seg.get("content", "")
                if not text:
                    continue
                run = para.add_run(text)
                run.font.name = self.font
                run.font.size = Pt(self.font_size)
                if seg.get("bold"):
                    run.font.bold = True
                if seg.get("italic"):
                    run.font.italic = True
                if seg.get("strikethrough"):
                    run.font.strike = True
                # 设置中文字体
                run.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

            elif seg_type == "code_inline":
                text = seg.get("content", "")
                run = para.add_run(text)
                run.font.name = 'Courier New'
                run.font.size = Pt(self.font_size - 1)
                # 浅灰底纹通过字符底纹实现
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'clear')
                shading.set(qn('w:color'), 'auto')
                shading.set(qn('w:fill'), 'F0F0F0')
                run.element.rPr.append(shading)

            elif seg_type == "link":
                text = seg.get("content", "")
                href = seg.get("href", "")
                if href:
                    self._add_hyperlink(para, text, href)
                else:
                    run = para.add_run(text)
                    run.font.name = self.font
                    run.font.size = Pt(self.font_size)

            elif seg_type == "image":
                # 内联图片（较少见）
                pass

            elif seg_type == "footnote_ref":
                label = seg.get("label", "")
                self._footnote_counter += 1
                run = para.add_run(f'[{self._footnote_counter}]')
                run.font.size = Pt(8)
                run.font.superscript = True
                run.font.color.rgb = RGBColor(0, 102, 204)

            elif seg_type == "softbreak":
                para.add_run('\n')

            elif seg_type == "hardbreak":
                para.add_run('\n')

    def _add_hyperlink(self, para, text: str, url: str):
        """
        在段落中添加超链接

        Args:
            para: 段落对象
            text: 链接显示文本
            url: 链接URL
        """
        # 获取文档关系部分
        part = para.part
        r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)

        # 创建超链接元素
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)

        # 创建run
        new_run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')

        # 蓝色字体
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '0563C1')
        rPr.append(color)

        # 下划线
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

        # 字体设置
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), self.font)
        rFonts.set(qn('w:eastAsia'), self.font)
        rPr.append(rFonts)

        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(self.font_size * 2)))  # half-point
        rPr.append(sz)

        new_run.append(rPr)

        # 文本内容
        t = OxmlElement('w:t')
        t.text = text
        new_run.append(t)

        hyperlink.append(new_run)
        para._p.append(hyperlink)

    def _add_list(self, node: Dict[str, Any], level: int = 0):
        """
        添加列表

        Args:
            node: 列表节点
            level: 列表级别
        """
        children = node.get('children', [])
        ordered = node.get('attributes', {}).get('ordered', False)

        for item in children:
            if item.get('type') == 'list_item':
                # 添加列表项
                if ordered:
                    para = self.doc.add_paragraph(style='List Number')
                else:
                    para = self.doc.add_paragraph(style='List Bullet')

                # 设置缩进
                para.paragraph_format.left_indent = Cm(1.27 * level)

                # 获取列表项中的段落内容（支持内联格式）
                item_paragraphs = [c for c in item.get('children', []) if c.get('type') == 'paragraph']
                if item_paragraphs:
                    first_para = item_paragraphs[0]
                    segments = first_para.get('children', [])
                    content = first_para.get('content', '')

                    # 任务列表：添加复选框前缀
                    task_checked = first_para.get('attributes', {}).get('task_checked')
                    if task_checked is not None:
                        checkbox = '☑ ' if task_checked else '☐ '
                        cb_run = para.add_run(checkbox)
                        cb_run.font.name = self.font
                        cb_run.font.size = Pt(self.font_size)

                    self._process_inline_content(para, content, segments if segments else None)
                else:
                    item_content = self._get_list_item_content(item)
                    run = para.add_run(item_content)

                # 设置中文字体
                for run in para.runs:
                    if not run.font.name:
                        run.font.name = self.font
                        run.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

                # 处理嵌套列表
                for child in item.get('children', []):
                    if child.get('type') == 'list':
                        self._add_list(child, level + 1)

    def _get_list_item_content(self, node: Dict[str, Any]) -> str:
        """
        获取列表项内容

        Args:
            node: 列表项节点

        Returns:
            内容文本
        """
        content_parts = []
        for child in node.get('children', []):
            if child.get('type') == 'paragraph':
                content_parts.append(child.get('content', ''))
        return ' '.join(content_parts)

    def _add_table(self, node: Dict[str, Any]):
        """
        添加表格

        Args:
            node: 表格节点
        """
        children = node.get('children', [])

        if not children:
            return

        # 计算行列数
        rows = len(children)
        cols = len(children[0].get('children', [])) if children else 0

        if rows == 0 or cols == 0:
            return

        # 创建表格
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'

        # 填充表格内容
        for i, row_node in enumerate(children):
            row_cells = row_node.get('children', [])
            for j, cell_node in enumerate(row_cells):
                if j < cols:
                    cell = table.cell(i, j)
                    cell.text = cell_node.get('content', '')

                    # 设置中文字体和对齐
                    align = cell_node.get('attributes', {}).get('align', '')
                    for para in cell.paragraphs:
                        if align == 'center':
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        elif align == 'right':
                            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        for run in para.runs:
                            run.font.name = self.font
                            run.font.size = Pt(self.font_size)
                            run.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)

    def _add_code_block(self, node: Dict[str, Any]):
        """
        添加代码块

        Args:
            node: 代码块节点
        """
        content = node.get('content', '')
        language = node.get('attributes', {}).get('language', '')

        # 添加代码块段落
        para = self.doc.add_paragraph()

        # 设置代码块样式
        para.paragraph_format.left_indent = Cm(1)
        para.paragraph_format.right_indent = Cm(1)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(6)

        # 添加代码内容
        run = para.add_run(content)
        run.font.name = 'Courier New'
        run.font.size = Pt(10)

        # 设置背景色（浅灰色）
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F5F5F5')
        pPr.append(shd)

        # 如果有语言标识，添加语言标签
        if language:
            lang_para = self.doc.add_paragraph()
            lang_run = lang_para.add_run(f'Language: {language}')
            lang_run.font.size = Pt(8)
            lang_run.font.color.rgb = RGBColor(128, 128, 128)

    def _add_flowchart(self, node: Dict[str, Any], chart_type: str):
        """
        添加流程图

        Args:
            node: 代码块节点
            chart_type: 流程图类型（mermaid, plantuml）
        """
        content = node.get('content', '')

        # 生成输出路径
        self.flowchart_counter += 1
        output_dir = Path(self.output_dir) if self.output_dir else Path('.')
        output_dir.mkdir(parents=True, exist_ok=True)
        image_path = output_dir / f'flowchart_{self.flowchart_counter}.png'

        # 渲染流程图
        success = self.flowchart_processor.render_flowchart(
            content, chart_type, str(image_path)
        )

        if success and image_path.exists():
            # 添加流程图标题
            title_para = self.doc.add_paragraph()
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_run = title_para.add_run(f'流程图 {self.flowchart_counter}')
            title_run.font.size = Pt(9)
            title_run.font.color.rgb = RGBColor(128, 128, 128)

            # 插入图片
            self.doc.add_picture(str(image_path), width=Inches(5))

            # 添加图片说明
            caption_para = self.doc.add_paragraph()
            caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption_para.add_run(f'图 {self.flowchart_counter}')
            caption_run.font.size = Pt(9)
            caption_run.font.color.rgb = RGBColor(128, 128, 128)
        else:
            # 渲染失败，添加代码块
            self.log(f"流程图渲染失败，添加为代码块", "warning")
            self._add_code_block(node)

    def _add_blockquote(self, node: Dict[str, Any]):
        """
        添加引用

        Args:
            node: 引用节点
        """
        children = node.get('children', [])

        for child in children:
            if child.get('type') == 'paragraph':
                # 添加引用段落
                para = self.doc.add_paragraph()

                # 设置引用样式
                para.paragraph_format.left_indent = Cm(2)
                para.paragraph_format.right_indent = Cm(2)

                # 添加引用内容（支持内联格式）
                content = child.get('content', '')
                segments = child.get('children', [])
                if segments:
                    self._process_inline_content(para, content, segments)
                else:
                    run = para.add_run(content)
                    run.font.name = self.font
                    run.font.size = Pt(self.font_size)
                    run.font.italic = True

                # 设置左边框（灰色竖线）
                pPr = para._p.get_or_add_pPr()
                pBdr = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                left.set(qn('w:val'), 'single')
                left.set(qn('w:sz'), '18')
                left.set(qn('w:space'), '4')
                left.set(qn('w:color'), 'BFBFBF')
                pBdr.append(left)
                pPr.append(pBdr)

    def _add_image(self, node: Dict[str, Any]):
        """
        添加图片

        Args:
            node: 图片节点
        """
        src = node.get('attributes', {}).get('src', '')
        alt = node.get('attributes', {}).get('alt', '')

        if not src:
            self.log(f"图片路径为空: {alt}", "warning")
            return

        try:
            image_path = None

            if src.startswith('http://') or src.startswith('https://'):
                # 远程图片，下载到临时目录
                image_path = self._download_remote_image(src)
                if not image_path:
                    return
            else:
                # 本地图片
                p = Path(src)
                if p.exists():
                    image_path = p
                elif self.input_dir:
                    # 尝试相对于输入文件的路径
                    p = self.input_dir / src
                    if p.exists():
                        image_path = p

                if not image_path:
                    self.log(f"图片不存在: {src}", "warning")
                    return

            # 插入图片
            self.doc.add_picture(str(image_path), width=Inches(4))

            # 添加图片标题
            if alt:
                caption = self.doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(alt)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(128, 128, 128)

        except Exception as e:
            self.log(f"插入图片失败: {e}", "error")

    def _download_remote_image(self, url: str) -> Path:
        """
        下载远程图片到临时目录

        Args:
            url: 图片URL

        Returns:
            临时文件路径，失败返回None
        """
        try:
            # 从URL推断文件扩展名
            ext = Path(url.split('?')[0]).suffix or '.png'
            # 创建临时文件
            tmp = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
            tmp_path = Path(tmp.name)
            # 下载
            self.log(f"下载远程图片: {url}")
            with urlopen(url, timeout=15) as resp:
                tmp.write(resp.read())
            tmp.close()
            return tmp_path
        except (URLError, OSError, TimeoutError) as e:
            self.log(f"远程图片下载失败: {url} ({e})", "warning")
            return None

    def _add_hr(self, node: Dict[str, Any]):
        """
        添加分割线

        Args:
            node: 分割线节点（用于接口一致性）
        """
        para = self.doc.add_paragraph()

        # 通过段落底部边框实现水平分割线
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'BFBFBF')
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_footnote_block(self, node: Dict[str, Any]):
        """
        添加脚注块（文档末尾的脚注列表）

        Args:
            node: 脚注块节点
        """
        children = node.get('children', [])
        if not children:
            return

        # 添加脚注分隔线
        sep_para = self.doc.add_paragraph()
        pPr = sep_para._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'BFBFBF')
        pBdr.append(bottom)
        pPr.append(pBdr)

        for i, footnote in enumerate(children):
            if footnote.get('type') != 'footnote':
                continue
            label = footnote.get('attributes', {}).get('label', '')
            fn_children = footnote.get('children', [])

            para = self.doc.add_paragraph()
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)

            # 脚注编号
            num_run = para.add_run(f'[{i + 1}] ')
            num_run.font.size = Pt(9)
            num_run.font.superscript = True
            num_run.font.color.rgb = RGBColor(0, 102, 204)

            # 脚注内容
            for fn_child in fn_children:
                if fn_child.get('type') == 'paragraph':
                    content = fn_child.get('content', '')
                    segments = fn_child.get('children', [])
                    if segments:
                        self._process_inline_content(para, content, segments)
                    else:
                        run = para.add_run(content)
                        run.font.name = self.font
                        run.font.size = Pt(9)